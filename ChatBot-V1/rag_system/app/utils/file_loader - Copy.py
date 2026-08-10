import json
import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


def load_file(file_path: str) -> Optional[str]:
    path = Path(file_path)
    if not path.exists():
        logger.error(f"File not found: {file_path}")
        return None
    suffix = path.suffix.lower()
    loaders = {
        ".txt":  _load_txt,
        ".pdf":  _load_pdf,
        ".docx": _load_docx,
        ".csv":  _load_csv,
        ".xlsx": _load_xlsx,
        ".json": _load_json,
    }
    loader = loaders.get(suffix)
    if not loader:
        logger.warning(f"Unsupported file type: {suffix}")
        return None
    try:
        text = loader(path)
        logger.info(f"Loaded {file_path} ({len(text)} chars)")
        return text
    except Exception as e:
        logger.error(f"Failed to load {file_path}: {e}", exc_info=True)
        return None


def _load_txt(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="latin-1")


def _load_pdf(path: Path) -> str:
    import PyPDF2
    texts = []
    with open(path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for i, page in enumerate(reader.pages):
            try:
                texts.append(page.extract_text() or "")
            except Exception as e:
                logger.warning(f"Skipping PDF page {i}: {e}")
    return "\n\n".join(texts)


def _load_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


def _load_csv(path: Path) -> str:
    import pandas as pd
    df = pd.read_csv(path)
    rows = []
    for _, row in df.iterrows():
        parts = [f"{col}: {val}" for col, val in row.items()
                 if str(val).strip() not in ("", "nan", "None")]
        if parts:
            rows.append(". ".join(parts))
    return "\n".join(rows)


def _load_xlsx(path: Path) -> str:
    import pandas as pd
    sheets = pd.read_excel(path, sheet_name=None)
    out = []
    for name, df in sheets.items():
        out.append(f"=== Sheet: {name} ===")
        for _, row in df.iterrows():
            parts = [f"{col}: {val}" for col, val in row.items()
                     if str(val).strip() not in ("", "nan", "None")]
            if parts:
                out.append(". ".join(parts))
    return "\n".join(out)


def _load_json(path: Path) -> str:
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    def flatten(obj, prefix=""):
        lines = []
        if isinstance(obj, dict):
            for k, v in obj.items():
                lines.append(flatten(v, f"{prefix}.{k}" if prefix else k))
        elif isinstance(obj, list):
            for i, v in enumerate(obj):
                lines.append(flatten(v, f"{prefix}[{i}]"))
        else:
            if str(obj).strip():
                lines.append(f"{prefix}: {obj}")
        return "\n".join(filter(None, lines))
    return flatten(data)
